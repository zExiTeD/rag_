import os
import fitz
import docx
import whisper
from pptx import Presentation

# --- Supported extensions ---
TEXT_EXTS = {".txt", ".docx", ".pptx", ".pdf"}
AUDIO_EXTS = {".mp3", ".wav", ".flac", ".aac", ".ogg", ".m4a"}


def extract_audio(file_path, model_name="base"):
    """
    Extracts speech from audio using OpenAI Whisper.
    Returns dict with "text" and "images" keys.
    """
    ext = os.path.splitext(file_path)[1].lower()
    if ext not in AUDIO_EXTS:
        raise ValueError(f"Unsupported audio type: {ext}")

    file_name = os.path.basename(file_path)
    model = whisper.load_model(model_name)
    result = model.transcribe(file_path)

    text_segments = []
    for seg_id, seg in enumerate(result["segments"], start=1):
        text_segments.append({
            "file": file_name,
            "page": None,
            "line": None,
            "segment_id": seg_id,
            "start": seg["start"],
            "end": seg["end"],
            "content": seg["text"].strip()
        })

    return {"text": text_segments, "images": []}


def extract_content_text(file_path):
    """
    Extracts text and images from docx, pptx, txt, and pdf files.
    Returns dict with "text" and "images" keys.
    """
    ext = os.path.splitext(file_path)[1].lower()
    result = {"text": [], "images": []}
    file_name = os.path.basename(file_path)

    if ext == ".txt":
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            for i, line in enumerate(f, start=1):
                if line.strip():
                    result["text"].append({
                        "file": file_name,
                        "page": 1,
                        "line": i,
                        "segment_id": None,
                        "start": None,
                        "end": None,
                        "content": line.strip()
                    })

    elif ext == ".docx":
        doc = docx.Document(file_path)
        for i, para in enumerate(doc.paragraphs, start=1):
            if para.text.strip():
                result["text"].append({
                    "file": file_name,
                    "page": 1,
                    "line": i,
                    "segment_id": None,
                    "start": None,
                    "end": None,
                    "content": para.text.strip()
                })
        # Images
        for rel in doc.part.rels.values():
            if "image" in rel.target_ref:
                image_data = rel.target_part.blob
                result["images"].append({
                    "file": file_name,
                    "page": 1,
                    "line": None,
                    "image_bytes": image_data
                })

    elif ext == ".pptx":
        prs = Presentation(file_path)
        for slide_num, slide in enumerate(prs.slides, start=1):
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    result["text"].append({
                        "file": file_name,
                        "page": slide_num,
                        "line": None,
                        "segment_id": None,
                        "start": None,
                        "end": None,
                        "content": shape.text.strip()
                    })
                if shape.shape_type == 13:  # Picture
                    image = shape.image
                    result["images"].append({
                        "file": file_name,
                        "page": slide_num,
                        "line": None,
                        "image_bytes": image.blob
                    })

    elif ext == ".pdf":
        doc = fitz.open(file_path)
        for page_num, page in enumerate(doc, start=1):
            text = page.get_text("text")
            if text:
                for line_num, line in enumerate(text.splitlines(), start=1):
                    if line.strip():
                        result["text"].append({
                            "file": file_name,
                            "page": page_num,
                            "line": line_num,
                            "segment_id": None,
                            "start": None,
                            "end": None,
                            "content": line.strip()
                        })
            # Images
            for img_index, img in enumerate(page.get_images(full=True)):
                xref = img[0]
                pix = fitz.Pixmap(doc, xref)
                if pix.n < 5:  # RGB or Gray
                    img_bytes = pix.tobytes("png")
                else:  # CMYK
                    pix = fitz.Pixmap(fitz.csRGB, pix)
                    img_bytes = pix.tobytes("png")
                result["images"].append({
                    "file": file_name,
                    "page": page_num,
                    "line": None,
                    "image_bytes": img_bytes
                })

    else:
        raise ValueError(f"Unsupported file type: {ext}")

    return result


def extract_content(file_path, model_name="base"):
    """
    Dispatcher: picks the right extractor depending on file type.
    """
    ext = os.path.splitext(file_path)[1].lower()

    if ext in AUDIO_EXTS:
        return extract_audio(file_path, model_name=model_name)
    elif ext in TEXT_EXTS:
        return extract_content_text(file_path)
    else:
        raise ValueError(f"Unsupported file type: {ext}")

